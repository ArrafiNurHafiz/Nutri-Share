import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
                <w:bottom w:val="single" w:sz="8" w:space="0" w:color="10B981"/>
                <w:left w:val="none"/>
                <w:right w:val="none"/>
                <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
                <w:insideV w:val="none"/>
            </w:tblBorders>
        ''')
        tblPr[0].append(borders)

def build_docx(output_path: str):
    doc = Document()

    # Page setup - Margins 1 inch (2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles & Colors
    c_primary = RGBColor(16, 185, 129)     # Emerald 500
    c_dark = RGBColor(15, 23, 42)          # Slate 900
    c_sub = RGBColor(71, 85, 105)          # Slate 600
    c_quote_bg = "F1F5F9"                  # Slate 100

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("NUTRISHARE")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = c_primary

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(16)
    run_sub = p_sub.add_run("Live Demonstration Script & Technical Presentation Manual\n(Focus: Real-Time Hybrid Entropy-TOPSIS & Exclusive Competitive Advantage)")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(13)
    run_sub.font.bold = True
    run_sub.font.color.rgb = c_dark

    # Meta box
    tbl_meta = doc.add_table(rows=1, cols=1)
    tbl_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_meta = tbl_meta.rows[0].cells[0]
    set_cell_background(cell_meta, "F8FAFC")
    set_cell_margins(cell_meta, top=140, bottom=140, left=200, right=200)
    p_meta = cell_meta.paragraphs[0]
    p_meta.paragraph_format.space_after = Pt(0)
    r_m = p_meta.add_run("📋 Presentation Blueprint & Executive Summary:\n"
                         "• Target Audience: Grand Jury, Academic Evaluators, & Food-Tech / Social Impact Stakeholders\n"
                         "• Recommended Duration: 6 – 8 Minutes (Live Demonstration) + 5 Minutes (Technical Q&A Defense)\n"
                         "• System Architecture: Full-Stack React 19 + Asynchronous FastAPI (Python 3.14) + Leaflet GIS + SSE\n"
                         "• Core Thesis: Transforming surplus food rescue from an unfair 'First-Come, First-Served' race into an intelligent, scientifically-optimized Decision Support System based on National AKG Nutritional Standards.")
    r_m.font.name = "Arial"
    r_m.font.size = Pt(9.5)
    r_m.font.color.rgb = c_sub

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 1: Competitive Matrix
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    r_h1 = h1.add_run("1. Competitive Advantage Matrix (NutriShare vs. Existing Platforms)")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(14)
    r_h1.font.bold = True
    r_h1.font.color.rgb = c_dark

    table_data = [
        ["Key Dimensions", "Conventional Food Rescue Apps\n(Surplus, Olio, Too Good To Go)", "NutriShare Platform (Our System)", "Value Added / Impact"],
        ["Allocation Mechanism", "First-Come, First-Served (FCFS) / Fastest click wins.", "Multi-Criteria Hybrid Entropy-TOPSIS Decision Support System.", "Eliminates monopoly by high-speed internet users; allocates to shelters with the most acute nutritional deficit."],
        ["Weight Determination", "Static / Subjective / Fixed operator assumptions.", "Dynamic Shannon Entropy (wj) combined with Domain Nutrition Policy.", "Weights calculate objectively in real-time from actual recipient variance, eliminating human bias."],
        ["Nutritional Standards", "Crude food weight (kg) or generic broad categories.", "Indonesian National AKG Standard (Permenkes No. 28/2019).", "Calculates daily protein, calorie, and micronutrient deficits specifically tailored to shelter demographic intake."],
        ["Fairness & Anti-Monopoly", "None (nearby active shelters win repeatedly).", "Dynamic Fairness Penalty (Criterion C5).", "Automatically penalizes recent recipients so allocation seamlessly rotates across underserved institutions."],
        ["Shelf-Life Degradation", "Food spoils if initial claimant fails to collect.", "Tiered Shelf-Life Priority Escalation Mechanism.", "Automated priority window (1/4 of safe shelf life) cascades offers to Rank #2, #3 before food expires."],
        ["Food Safety Integrity", "Unverified peer-to-peer open forum.", "Closed-Loop Governance (KYC + Arrival Physical Inspection).", "Mandatory legal shelter license verification and donor hygiene certification prior to transaction authorization."],
        ["Macro ESG & SDG Analytics", "Static quarterly / annual PDF reports.", "Real-Time Ecological & Nutrition Metrics (CO2e, Protein, Portions).", "Live tracking of prevented greenhouse emissions and grams of protein delivered for verified SDG 2, 3, and 12 reporting."]
    ]

    tbl_comp = doc.add_table(rows=len(table_data), cols=4)
    tbl_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_comp)

    col_widths = [Inches(1.3), Inches(1.8), Inches(1.8), Inches(1.6)]

    for row_idx, row in enumerate(table_data):
        for col_idx, text in enumerate(row):
            cell = tbl_comp.rows[row_idx].cells[col_idx]
            cell.width = col_widths[col_idx]
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(text)
            run.font.name = "Arial"
            if row_idx == 0:
                set_cell_background(cell, "10B981")
                run.font.bold = True
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                if row_idx % 2 == 1:
                    set_cell_background(cell, "F8FAFC")
                else:
                    set_cell_background(cell, "FFFFFF")
                run.font.size = Pt(8.5)
                run.font.color.rgb = c_dark
                if col_idx == 2:
                    run.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 2: Mathematical Formula
    h2 = doc.add_heading(level=1)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)
    r_h2 = h2.add_run("2. Mathematical Formulation: Real-Time Hybrid Entropy-TOPSIS")
    r_h2.font.name = "Arial"
    r_h2.font.size = Pt(14)
    r_h2.font.bold = True
    r_h2.font.color.rgb = c_dark

    p_topsis_desc = doc.add_paragraph()
    r_td = p_topsis_desc.add_run(
        "NutriShare's decision core executes a vectorized Hybrid Entropy-TOPSIS model in NumPy (<2ms for 50 recipients). The model evaluates 5 dynamic criteria:\n\n"
        "• C1 (Benefit): Protein Need Fulfillment Ratio\n"
        "  C1 = min(100, (Total Donation Protein / Recipient Remaining Deficit) × 100) × (1 - 0.5 × Fulfillment Ratio)\n\n"
        "• C2 (Benefit): Shelter Health Urgency\n"
        "  C2 = (Vulnerability Score of Toddlers & Elderly [×1000 if Disaster Emergency Active]) × (1 - 0.85 × Fulfillment Ratio)\n\n"
        "• C3 (Benefit): Food Freshness / Safe Shelf-Life Window\n"
        "  C3 = max((Valid Until Timestamp - Current Time) / 1 Hour, 0.1)\n\n"
        "• C4 (Cost): Haversine Geographical Distance\n"
        "  C4 = 2R · arcsin(√(sin²(Δlat/2) + cos(lat1)·cos(lat2)·sin²(Δlon/2))) [in km]\n\n"
        "• C5 (Benefit): Distribution Fairness Index\n"
        "  C5 = Days Elapsed Since Last Donation Received [in Days]\n\n"
        "Objective Weight Calculation (Shannon Entropy):\n"
        "  P_ij = R_ij / Σ R_kj  -->  E_j = -k Σ P_ij ln(P_ij)  -->  d_j = 1 - E_j  -->  w_entropy = d_j / Σ d_j\n\n"
        "Hybrid Policy Blending:\n"
        "  w_j = (0.50 × w_policy) + (0.50 × w_entropy)\n"
        "where w_policy = [C1: 0.25, C2: 0.25, C3: 0.15, C4: 0.20, C5: 0.15] ensures domain integrity while entropy adapts to live data variance.\n\n"
        "TOPSIS Relative Closeness Score:\n"
        "  V_i = D_i^- / (D_i^+ + D_i^-)  [where D+ is distance to Ideal Positive and D- is distance to Ideal Negative]"
    )
    r_td.font.name = "Arial"
    r_td.font.size = Pt(9.0)
    r_td.font.color.rgb = c_dark

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 3: Live Demo Script
    h3 = doc.add_heading(level=1)
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)
    r_h3 = h3.add_run("3. End-to-End Live Demonstration Script (Step-by-Step)")
    r_h3.font.name = "Arial"
    r_h3.font.size = Pt(14)
    r_h3.font.bold = True
    r_h3.font.color.rgb = c_dark

    acts = [
        {
            "act": "ACT 1: THE PARADOX & PLATFORM HOOK (~45 Seconds)",
            "route": "Screen Route: http://localhost:5173/ (Landing Page)",
            "action": "Presenter Action: Open landing page. Scroll through Hero Section, showcase live impact counters (Surplus Rescued, Protein Distributed, Verified Shelters), and highlight the 3 Core Pillars.",
            "speech": "Good morning, distinguished members of the jury.\n\nEvery year, Indonesia loses between 23 to 48 million tons of edible food, generating an economic deficit exceeding 500 trillion Rupiah. Yet, right in our immediate neighborhood, orphanages and social shelters suffer from chronic hidden hunger and severe protein deficiencies.\n\nWhy haven't existing food rescue apps solved this crisis? Because they rely on a fundamentally flawed mechanism: 'First-Come, First-Served'. Whoever has the fastest internet and quickest thumbs grabs the food—regardless of whether they actually need it or whether the nutritional profile meets their dietary needs.\n\nThis is NutriShare — Indonesia's first intelligent, nutrition-optimized Decision Support System that transforms food surplus allocation from an unfair speed race into a scientifically verified, equity-driven distribution."
        },
        {
            "act": "ACT 2: DONOR WORKFLOW & NUTRITIONAL PROFILING (~60 Seconds)",
            "route": "Screen Route: http://localhost:5173/donor (Donor Dashboard)",
            "action": "Presenter Action: Log in as Hotel/Restaurant Donor. Click 'Tambah Donasi Surplus' (Create Food Listing), select preset 'Paket Nasi & Lauk Komplit', show pre-filled nutrition values (50 portions, 26g protein/portion, 4 hours shelf life), then click 'Publikasikan Donasi' (Publish).",
            "speech": "Let us examine the donor workflow—such as a banquet manager at a partner hotel.\n\nAt the conclusion of an event, donors can log surplus meals in under ten seconds using our intelligent presets. Crucially, donors specify the exact portion count, the safe consumption shelf-life window, and the macronutrient density—particularly protein content.\n\nThe instant I press 'Publish', NutriShare does NOT merely dump this listing onto an unmoderated public board. Our asynchronous backend immediately feeds this food profile into our scientific decision engine."
        },
        {
            "act": "ACT 3: THE ENGINE — REAL-TIME HYBRID ENTROPY-TOPSIS AUDIT (~90 Seconds)",
            "route": "Screen Route: http://localhost:5173/donor -> Click 'Audit Rekomendasi TOPSIS' (or via /admin)",
            "action": "Presenter Action: Open TOPSIS Transparency Audit Modal. Show the weighted decision matrix, 5 criteria calculations, real-time Shannon Entropy weights (wj), Euclidean distances (D+, D-), and the final Preference Scores (Vi).",
            "speech": "Here lies NutriShare's groundbreaking technological innovation: the Real-Time Hybrid Entropy-TOPSIS Decision Engine.\n\nWhy do we utilize a Hybrid approach?\nFirst, Shannon Entropy measures the statistical information dispersion across all eligible shelters in real-time, computing objective criteria weights (wj) free from human subjectivity.\nSecond, TOPSIS calculates the exact Euclidean distance of each shelter to the Ideal Positive Solution (A+) and Ideal Negative Solution (A-).\n\nAs audited live on screen, 'Panti Asuhan Kasih Ibu' claims Rank #1 with an optimal Relative Closeness Score of 0.892. This occurs because their children face an acute daily protein deficit, they house vulnerable toddlers, they are located within a close 3.2 km radius, and they have not received aid in the past 8 days.\n\nThis is true nutritional justice driven by empirical data, not random luck."
        },
        {
            "act": "ACT 4: RECIPIENT DASHBOARD, AKG GAUGE & PRIORITY CLAIM (~75 Seconds)",
            "route": "Screen Route: http://localhost:5173/recipient (Dashboard for Rank #1 Shelter)",
            "action": "Presenter Action: Switch to Rank #1 Shelter account. Demonstrate the AKG Nutritional Fulfillment Gauge, point out the instant Server-Sent Events (SSE) priority notification, show the countdown timer, and click 'Klaim Donasi' (Claim).",
            "speech": "Switching now to the recipient shelter's dashboard.\n\nShelter administrators monitor their live daily AKG Nutritional Gauge based on Indonesian Ministry of Health standards. When the donor published the surplus, our Server-Sent Events (SSE) immediately pushed an exclusive priority alert solely to this top-ranked shelter.\n\nThey receive a dedicated, protected time window to claim. If the shelter is unable to receive the food and the window expires, NutriShare's automated Tiered Shelf-Life Priority Escalation cascades the offer to Rank #2, ensuring zero food spoilage.\n\nI will now click 'Claim Donation'. The status transitions instantly into active transit."
        },
        {
            "act": "ACT 5: LEAFLET GEOSPATIAL TRACKING & HANDOVER INSPECTION (~45 Seconds)",
            "route": "Screen Route: http://localhost:5173/recipient -> Click 'Lacak Pengiriman' (Live Tracking)",
            "action": "Presenter Action: Open interactive Leaflet routing map. Show donor pickup coordinates, courier route, and shelter destination. Simulate arrival, click 'Konfirmasi Penerimaan', and submit physical condition verification & star rating.",
            "speech": "Once claimed, NutriShare coordinates logistics with real-time Leaflet GIS mapping, providing live coordinates, travel distance, and estimated arrival times.\n\nUpon delivery, the shelter conducts a mandatory physical inspection to verify food freshness before confirming handover in the app. This closes the operational loop: the shelter's AKG intake history updates instantly, and the donor is awarded social impact points and zero-waste badges."
        },
        {
            "act": "ACT 6: ADMIN GOVERNANCE & MACRO SDG ANALYTICS (~45 Seconds)",
            "route": "Screen Route: http://localhost:5173/admin (Admin Dashboard)",
            "action": "Presenter Action: Open Admin Dashboard. Show KYC Verification tab for institutional licenses and food safety certifications. Showcase macro analytics (Total kg rescued, grams of protein delivered, CO2e greenhouse emissions prevented).",
            "speech": "Platform integrity and food safety are governed through the Admin Dashboard.\n\nTo prevent fraud and guarantee food hygiene, every donor and recipient must pass rigorous administrative KYC verification before participating.\n\nFurthermore, administrators monitor macro-level ecological and social impacts—such as total protein grams delivered and metric tons of CO2e emissions prevented—delivering auditable, empirical reporting for UN Sustainable Development Goals 2, 3, and 12."
        },
        {
            "act": "ACT 7: CLOSING PITCH (~30 Seconds)",
            "route": "Screen Route: Return to http://localhost:5173/ (Landing Page)",
            "action": "Presenter Action: Conclude on the clean landing page.",
            "speech": "In conclusion, NutriShare elevates food surplus redistribution from an environmental burden into measurable social and nutritional equity.\n\nBy uniting the mathematical rigor of Hybrid Entropy-TOPSIS, national AKG standards, and end-to-end transparent governance, NutriShare provides the foundational digital infrastructure for the future of sustainable food security.\n\nThank you very much. We are now ready for your questions."
        }
    ]

    for item in acts:
        h_act = doc.add_heading(level=2)
        h_act.paragraph_format.space_before = Pt(12)
        h_act.paragraph_format.space_after = Pt(2)
        r_act = h_act.add_run(item["act"])
        r_act.font.name = "Arial"
        r_act.font.size = Pt(11)
        r_act.font.bold = True
        r_act.font.color.rgb = c_primary

        p_info = doc.add_paragraph()
        p_info.paragraph_format.space_after = Pt(4)
        r_info = p_info.add_run(f"📍 {item['route']}\n👉 {item['action']}")
        r_info.font.name = "Arial"
        r_info.font.size = Pt(8.5)
        r_info.font.italic = True
        r_info.font.color.rgb = c_sub

        tbl_sp = doc.add_table(rows=1, cols=1)
        tbl_sp.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell_sp = tbl_sp.rows[0].cells[0]
        set_cell_background(cell_sp, c_quote_bg)
        set_cell_margins(cell_sp, top=100, bottom=100, left=150, right=150)
        p_sp = cell_sp.paragraphs[0]
        p_sp.paragraph_format.space_after = Pt(0)
        p_sp.paragraph_format.line_spacing = 1.15

        r_en_lbl = p_sp.add_run("🎙️ Spoken English Script:\n")
        r_en_lbl.font.bold = True
        r_en_lbl.font.size = Pt(9.5)
        r_en_lbl.font.color.rgb = c_dark

        r_en_txt = p_sp.add_run(f"\"{item['speech']}\"")
        r_en_txt.font.size = Pt(9.0)
        r_en_txt.font.color.rgb = c_dark

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Section 4: Technical Defense Sheet
    h4 = doc.add_heading(level=1)
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(6)
    r_h4 = h4.add_run("4. Technical Q&A Defense Sheet (Jury Question Defense)")
    r_h4.font.name = "Arial"
    r_h4.font.size = Pt(14)
    r_h4.font.bold = True
    r_h4.font.color.rgb = c_dark

    qa_data = [
        ["Critical Jury Question", "Technical & Theoretical Defense"],
        [
            "1. Why use Hybrid Entropy-TOPSIS instead of standard TOPSIS, AHP, or SAW?",
            "• Flaw of Standard TOPSIS / AHP: Criteria weights are assigned subjectively by human operators, introducing personal bias and vulnerability to manipulation.\n• Advantage of Shannon Entropy: Calculates objective weights (wj) based on the actual variance and information entropy in real-time recipient data.\n• Hybrid Blending (50% AKG Domain Policy + 50% Shannon Entropy) guarantees mathematical objectivity while preserving nutritional domain rules."
        ],
        [
            "2. How does the system prevent a single shelter from monopolizing all donations?",
            "• Enforced via Criterion C5 (Fairness / Days Since Last Donation) and Fulfillment Ratio Penalties.\n• When a shelter receives food today, their C5 score drops toward 0 and their daily AKG quota fills up, triggering an automatic reduction of up to 70% in their Preference Score (Vi) on subsequent donations to allow other shelters to win."
        ],
        [
            "3. How do you guarantee food safety and prevent spoilage during transit?",
            "• Multi-layer safeguarding:\n  1) Criteria C3 (Shelf-Life) and C4 (Haversine Distance) restrict matches strictly to shelters within a safe reach window.\n  2) Tiered Priority Escalation: Dedicated claim window (1/4 of safe shelf life) cascades unclaimed food to Rank #2 before expiry.\n  3) Handover Verification: Mandatory physical inspection checklist on the app before final handover."
        ],
        [
            "4. What is the computational latency of the algorithm on the backend?",
            "• The TOPSIS calculation is fully vectorized using NumPy in asynchronous FastAPI (Python 3.14).\n• Benchmarked computational latency for 50 candidate shelters is under 2 milliseconds (~1.8 ms), ensuring instant zero-lag real-time response."
        ],
        [
            "5. What scientific foundation governs the nutritional intake calculations?",
            "• Directly derived from the Indonesian Ministry of Health Regulation (Permenkes RI No. 28/2019) on Recommended Dietary Allowances (AKG), adjusted dynamically for shelter resident count, child age brackets, and special vulnerability factors."
        ]
    ]

    tbl_qa = doc.add_table(rows=len(qa_data), cols=2)
    tbl_qa.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_qa)

    qa_col_widths = [Inches(2.5), Inches(4.0)]

    for row_idx, row in enumerate(qa_data):
        for col_idx, text in enumerate(row):
            cell = tbl_qa.rows[row_idx].cells[col_idx]
            cell.width = qa_col_widths[col_idx]
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(text)
            run.font.name = "Arial"
            if row_idx == 0:
                set_cell_background(cell, "10B981")
                run.font.bold = True
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                if row_idx % 2 == 1:
                    set_cell_background(cell, "F8FAFC")
                else:
                    set_cell_background(cell, "FFFFFF")
                run.font.size = Pt(8.5)
                run.font.color.rgb = c_dark
                if col_idx == 0:
                    run.font.bold = True

    doc.save(output_path)
    print(f"File successfully created at: {output_path}")

if __name__ == "__main__":
    build_docx("docs/NutriShare_Naskah_Demo_Presentasi.docx")
    build_docx("docs/NutriShare_Live_Demo_Script_English.docx")
